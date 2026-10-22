"""从本机已获取的真实原始文件切出测试用的小样本（intake-02）。

测试不联网、也不依赖本机缓存：``tests/fixtures/intake/cache/`` 里放的是真实文件的**子集**（芝加哥都会区 16980），格式与原文件相同，
按同一个请求缓存键存放，所以测试走的是与真实构建完全相同的读取、标准化、映射与检查代码。每份样本的 ``.meta.json`` 记着原文件的
SHA-256、大小与切法；``sha256`` 是样本自己的。没有复制、排列或生成记录：每一行都原样来自原文件。

样本只为检验流程，**分布不代表现实**（例如住房只留了几个 PUMA 的一部分租住户）；测试不得据此判断统计结论。

    uv run python scripts/intake_fixtures.py [--cache data/raw/intake] [--out tests/fixtures/intake]

刻意不放进样本的请求（测试里它们是“没有取得”）：印第安纳与威斯康星的 PUMS（芝加哥都会区跨三个州，只留伊利诺伊）、HUD 公平
市场租金（真实获取时就被拒绝，``_fetch_log.json`` 抄了那次真实结果）。
"""

from __future__ import annotations

import argparse
import csv
import datetime as _dt
import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import yaml

import _bootstrap  # noqa: F401
from product.realdata import geo, sources
from product.realdata.acs import REP_WEIGHTS
from product.realdata.activities import PRIVATE_ACCESS
from product.realdata.background import AP_ENERGY
from product.realdata.config import DEFAULT_CONFIG, load_config
from product.realdata.fetch import DEFAULT_CACHE, Request, read_cached

CBSA = "16980"
PUMS_STATE = "17"
N_PUMAS = 4
N_OTHER_RENTERS = 40
N_INSTITUTIONS = 8
FIXED_TIME = (2024, 1, 1, 0, 0, 0)


def fixture_config(base: Path) -> dict:
    doc = load_config(base).model_dump(mode="json")
    doc["build"]["name"] = "us-v1-fixture-chicago"
    doc["metros"] = [m for m in doc["metros"] if m["cbsa"] == CBSA]
    doc["menus"]["years"] = [1, 2]
    return doc


# --------------------------------------------------------------------------- 容器格式


def rezip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in members.items():
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, data)
    return buf.getvalue()


def xlsx(sheets: dict[str, list[tuple]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(list(r))
    wb.properties.created = _dt.datetime(2024, 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def lines_keep(content: bytes, keep) -> bytes:
    """保留表头与 ``keep(行)`` 为真的行。按 latin-1 逐字节往返，所以保留的行与原文件逐字节相同（含 BOM 与换行）。"""
    lines = content.decode("latin-1").splitlines(keepends=True)
    return "".join([lines[0]] + [ln for ln in lines[1:] if ln.strip() and keep(ln)]).encode(
        "latin-1"
    )


# --------------------------------------------------------------------------- 各来源的切法


def cut_delineation(content: bytes) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    i = rows[2].index("CBSA Code")
    keep = rows[:3] + [r for r in rows[3:] if r[i] is not None and str(r[i]) == CBSA]
    return xlsx({ws.title: keep})


def touching_pumas(tract_puma: bytes, counties: set[str]) -> set[tuple[str, str]]:
    lines = tract_puma.decode("utf-8-sig").splitlines()
    h = [x.strip() for x in lines[0].split(",")]
    ist, ico, ipu = h.index("STATEFP"), h.index("COUNTYFP"), h.index("PUMA5CE")
    rows = [ln.split(",") for ln in lines[1:] if ln.strip()]
    return {(p[ist], p[ipu]) for p in rows if p[ist] + p[ico] in counties}


def cut_tract_puma(content: bytes, pumas: set[tuple[str, str]]) -> bytes:
    h = [x.strip() for x in content.decode("utf-8-sig").splitlines()[0].split(",")]
    ist, ipu = h.index("STATEFP"), h.index("PUMA5CE")

    def keep(ln: str) -> bool:
        p = ln.rstrip("\r\n").split(",")
        return (p[ist], p[ipu]) in pumas

    return lines_keep(content, keep)


def cut_gazetteer(content: bytes) -> bytes:
    z = zipfile.ZipFile(io.BytesIO(content))
    name = next(n for n in z.namelist() if n.endswith(".txt"))
    data = z.read(name)
    return rezip({name: lines_keep(data, lambda ln: ln.split("\t")[0].strip() == CBSA)})


def cut_geos(content: bytes, pumas: set[tuple[str, str]]) -> tuple[bytes, str]:
    geos = geo.parse_acs_geos(content)
    keep_ids = {
        k
        for k, g in geos.items()
        if (g["sumlevel"] == "310" and g["cbsa"] == CBSA)
        or (g["sumlevel"] == "795" and (g["state"], g["puma"]) in pumas)
    }
    cbsa_id = next(k for k, g in geos.items() if g["sumlevel"] == "310" and g["cbsa"] == CBSA)
    header = content.decode("latin-1").splitlines()[0].split("|")
    i = header.index("GEO_ID")
    return lines_keep(content, lambda ln: ln.split("|")[i] in keep_ids), cbsa_id


def cut_dictionary(content: bytes, variables: set[str]) -> bytes:
    def keep(ln: str) -> bool:
        row = next(csv.reader([ln]))
        return len(row) > 1 and row[1] in variables

    return lines_keep(content, keep)


def cut_pums(
    h_content: bytes,
    p_content: bytes,
    pumas: list[str],
    hvars: tuple[str, ...],
    pvars: tuple[str, ...],
) -> tuple[bytes, bytes, str]:
    zh = zipfile.ZipFile(io.BytesIO(h_content))
    hname = next(n for n in zh.namelist() if n.lower().startswith("psam_h") and n.endswith(".csv"))
    cols = list(dict.fromkeys([*hvars, *REP_WEIGHTS]))
    h = pd.read_csv(zh.open(hname), dtype=str, usecols=cols, keep_default_na=False)[cols]
    h = h[h["PUMA"].str.zfill(5).isin(pumas)]
    rent = pd.to_numeric(h["RNTP"], errors="coerce")
    renter_or_vacant = ((h["TEN"] == "3") | (h["VACS"] == "1")) & (rent > 0)
    market = renter_or_vacant & ((h["VACS"] == "1") | ((h["TEN"] == "3") & (h["MV"] == "1")))
    others = h[renter_or_vacant & ~market].groupby("PUMA", sort=True).head(N_OTHER_RENTERS)
    kept = pd.concat([h[market], others]).sort_index()
    zp = zipfile.ZipFile(io.BytesIO(p_content))
    pname = next(n for n in zp.namelist() if n.lower().startswith("psam_p") and n.endswith(".csv"))
    p = pd.read_csv(zp.open(pname), dtype=str, usecols=list(pvars), keep_default_na=False)[
        list(pvars)
    ]
    p = p[p["SERIALNO"].isin(set(kept["SERIALNO"]))]
    rule = (
        f"PUMA {', '.join(pumas)}（完整在芝加哥都会区内、编号最小的 {N_PUMAS} 个）：全部空置待租与 12 个月内迁入的租住户，"
        f"加上每个 PUMA 前 {N_OTHER_RENTERS} 户其他租住户；只留配置用到的列与 80 个重复权重；人员文件留这些住户的所有人"
    )
    return (
        rezip({hname: kept.to_csv(index=False, lineterminator="\n").encode("utf-8")}),
        rezip({pname: p.to_csv(index=False, lineterminator="\n").encode("utf-8")}),
        rule,
    )


def cut_oews(content: bytes) -> tuple[bytes, set[str]]:
    z = zipfile.ZipFile(io.BytesIO(content))
    name = next(n for n in z.namelist() if n.endswith("MSA_M2024_dl.xlsx"))
    df = pd.read_excel(z.open(name), dtype=str)
    df = df[df["AREA"] == CBSA]
    groups = df["O_GROUP"].str.lower()
    detailed = df[groups == "detailed"].sort_values("OCC_CODE")
    marked = detailed[
        detailed.apply(lambda r: any(str(v).strip() in ("#", "*", "**") for v in r.values), axis=1)
    ]
    keep = pd.concat([df[groups.isin(["total", "major"])], detailed.head(60), marked.head(10)])
    keep = keep[~keep.index.duplicated()].sort_index()
    rows = [tuple(df.columns)] + [
        tuple(None if pd.isna(v) else v for v in r) for r in keep.itertuples(index=False)
    ]
    return rezip({name: xlsx({"All May 2024 data": rows})}), set(keep["OCC_CODE"])


def cut_cpi(content: bytes) -> bytes:
    def keep(ln: str) -> bool:
        p = [x.strip() for x in ln.split("\t")]
        return len(p) > 2 and p[0] == "CUUR0000SA0" and p[1] >= "2022"

    return lines_keep(content, keep)


def basket_items(cfg_doc: dict) -> set[str]:
    out = set()
    for spec in cfg_doc["mapping"]["a_food"]["baskets"].values():
        out |= {str(line["item"]) for line in spec["items"]}
    return out


def cut_ap(series: bytes, data: bytes, items: set[str], areas: set[str]) -> tuple[bytes, bytes]:
    keep_series = set()

    def keep_s(ln: str) -> bool:
        p = [x.strip() for x in ln.split("\t")]
        if len(p) > 2 and p[1] in areas and p[2] in items:
            keep_series.add(p[0])
            return True
        return False

    s = lines_keep(series, keep_s)

    def keep_d(ln: str) -> bool:
        p = [x.strip() for x in ln.split("\t")]
        return len(p) > 1 and p[0] in keep_series and p[1] >= "2023"

    return s, lines_keep(data, keep_d)


def cut_ipeds(hd: bytes, py: bytes) -> tuple[bytes, bytes, set[str], set[str]]:
    zh, zp = zipfile.ZipFile(io.BytesIO(hd)), zipfile.ZipFile(io.BytesIO(py))
    hname = next(n for n in zh.namelist() if n.lower().endswith(".csv"))
    pname = next(n for n in zp.namelist() if n.lower().endswith(".csv"))
    hdf = pd.read_csv(zh.open(hname), dtype=str, encoding="latin-1", index_col=False)
    pdf = pd.read_csv(zp.open(pname), dtype=str, encoding="latin-1", index_col=False)
    hdf.columns = [c.strip().lstrip("﻿").lstrip("ï»¿") for c in hdf.columns]
    pdf.columns = [c.strip().lstrip("﻿").lstrip("ï»¿") for c in pdf.columns]
    local = set(hdf.loc[hdf["CBSA"] == CBSA, "UNITID"])
    with_programs = sorted(local & set(pdf["UNITID"]), key=int)[:N_INSTITUTIONS]
    ids = set(with_programs)
    cips = {
        str(v).strip()
        for _, r in pdf[pdf["UNITID"].isin(ids)].iterrows()
        for k in range(1, 7)
        if (v := r.get(f"CIPCODE{k}")) is not None and str(v).strip() not in ("", "nan", "A", "R")
    }

    def first_field(ln: str) -> str:
        return ln.split(",", 1)[0].strip().strip('"')

    hd_out = rezip({hname: lines_keep(zh.read(hname), lambda ln: first_field(ln) in ids)})
    py_out = rezip({pname: lines_keep(zp.read(pname), lambda ln: first_field(ln) in ids)})
    return hd_out, py_out, ids, cips


def cut_cip_soc(content: bytes, cips: set[str]) -> bytes:
    df = pd.read_excel(io.BytesIO(content), sheet_name="CIP-SOC", dtype=str)
    keep = df[df["CIP2020Code"].str.strip().isin(cips)]
    rows = [tuple(df.columns)] + [
        tuple(None if pd.isna(v) else v for v in r) for r in keep.itertuples(index=False)
    ]
    return xlsx({"CIP-SOC": rows})


def cut_aqi(content: bytes) -> bytes:
    z = zipfile.ZipFile(io.BytesIO(content))
    name = z.namelist()[0]
    header = next(csv.reader([z.read(name).decode("latin-1").splitlines()[0]]))
    i = header.index("CBSA Code")
    return rezip({name: lines_keep(z.read(name), lambda ln: next(csv.reader([ln]))[i] == CBSA)})


def cut_rpp(content: bytes) -> bytes:
    z = zipfile.ZipFile(io.BytesIO(content))
    name = next(n for n in z.namelist() if n.startswith("MARPP_MSA") and n.endswith(".csv"))

    def keep(ln: str) -> bool:
        row = next(csv.reader([ln]))
        return bool(row) and row[0].strip().strip('"') in ("00000", CBSA)

    return rezip({name: lines_keep(z.read(name), keep)})


def cut_summary_table(content: bytes, cbsa_geo_id: str) -> bytes:
    return lines_keep(content, lambda ln: ln.split("|")[0] == cbsa_geo_id)


def cut_chicago(content: bytes) -> tuple[bytes, str]:
    rows = sorted(json.loads(content), key=lambda r: str(r.get("activity_id")))

    def single_day(r: dict) -> bool:
        return (r.get("start_date") or "")[:10] == (r.get("end_date") or "")[:10] != ""

    def has_fee(r: dict) -> bool:
        return str(r.get("fee") or "").strip() != ""

    picks: list[dict] = []
    for rule, n in (
        (lambda r: single_day(r) and has_fee(r) and " to " in str(r.get("date_notes") or ""), 15),
        (lambda r: not single_day(r) and has_fee(r), 15),
        (lambda r: not has_fee(r), 5),
        (lambda r: r.get("event_cancelled") == "Y", 5),
        (lambda r: "&amp;" in json.dumps(r), 5),
    ):
        picks += [r for r in rows if rule(r) and r not in picks][:n]
    picks.sort(key=lambda r: str(r.get("activity_id")))
    rule_text = "按活动编号排序后：单日有价且写了时段的前 15 个、多日有价的前 15 个、没写价格的前 5 个、已取消的前 5 个、文字带 HTML 实体的前 5 个"
    return json.dumps(picks, ensure_ascii=False).encode("utf-8"), rule_text


def cut_osm(content: bytes, kind: str) -> bytes:
    doc = json.loads(content)
    els = [e for e in doc.get("elements", []) if e.get("tags")]
    if kind == "food":
        channel = [
            e for e in els if any(k in e["tags"] for k in ("takeaway", "delivery", "drive_through"))
        ]
        picks = els[:20] + [e for e in channel if e not in els[:20]][:10]
    else:
        private = [e for e in els if e["tags"].get("access") in PRIVATE_ACCESS]
        picks = els[:30] + [e for e in private if e not in els[:30]][:3]
    return json.dumps(
        {**{k: v for k, v in doc.items() if k != "elements"}, "elements": picks}, ensure_ascii=False
    ).encode("utf-8")


# --------------------------------------------------------------------------- 写样本


def put(out_cache: Path, req: Request, content: bytes, original, rule: str) -> int:
    base = out_cache / req.source_id / req.key[:16]
    base.parent.mkdir(parents=True, exist_ok=True)
    base.with_suffix(".bin").write_bytes(content)
    import hashlib

    meta = {
        "source_id": req.source_id,
        "url": req.display_url(),
        "method": req.method,
        "http_status": original.record.http_status,
        "sha256": hashlib.sha256(content).hexdigest(),
        "bytes": len(content),
        "retrieved_at": original.record.retrieved_at,
        "subset_of": {"sha256": original.record.raw_sha256, "bytes": original.record.bytes},
        "subset_rule": rule,
    }
    base.with_suffix(".meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n"
    )
    return len(content)


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--cache", default=str(DEFAULT_CACHE))
    ap.add_argument("--out", default="tests/fixtures/intake")
    ap.add_argument("--config", default=str(DEFAULT_CONFIG))
    args = ap.parse_args()
    cache, out = Path(args.cache), Path(args.out)
    out_cache = out / "cache"
    cfg_doc = fixture_config(Path(args.config))
    (out / "config_chicago.yaml").parent.mkdir(parents=True, exist_ok=True)
    (out / "config_chicago.yaml").write_text(
        "# 由 scripts/intake_fixtures.py 从 configs/intake/us_v1.yaml 派生：只留芝加哥都会区、菜单两年。不要手改。\n"
        + yaml.safe_dump(cfg_doc, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
        newline="\n",
    )
    cfg = load_config(out / "config_chicago.yaml")

    def real(req: Request):
        r = read_cached(req, cache)
        if not r.ok:
            raise SystemExit(
                f"本机缓存里没有 {req.display_url()}：先运行 scripts/intake_data.py fetch"
            )
        return r

    delin = real(geo.delineation_request())
    counties = {c.fips for c in geo.parse_delineation(delin.content)[CBSA].counties}
    states = sorted({c[:2] for c in counties})
    reqs = sources.all_requests(cfg, states)
    tp = real(reqs["tract_puma"])
    pumas = touching_pumas(tp.content, counties)
    sizes: dict[str, int] = {}

    def save(name: str, content: bytes, rule: str) -> None:
        sizes[name] = put(out_cache, reqs[name], content, real(reqs[name]), rule)

    save("cbsa_delineation", cut_delineation(delin.content), f"表头三行 + CBSA Code = {CBSA} 的县")
    save(
        "tract_puma",
        cut_tract_puma(tp.content, pumas),
        "与芝加哥都会区的县有交集的 PUMA 的全部普查区（含都会区外的普查区，保证完整 / 部分判定不变）",
    )
    save("gazetteer", cut_gazetteer(real(reqs["gazetteer"]).content), f"表头 + GEOID = {CBSA}")
    geos_bytes, cbsa_geo_id = cut_geos(real(reqs["acs_geos"]).content, pumas)
    save("acs_geos", geos_bytes, "CBSA 16980 与上述 PUMA 的地理行")
    variables = set(cfg.pums.housing_vars) | set(cfg.pums.person_vars)
    save(
        "pums_dictionary",
        cut_dictionary(real(reqs["pums_dictionary"]).content, variables),
        "配置用到的变量的 NAME / VAL 行",
    )
    g = geo.metro_geography(
        cfg.metros,
        delin.content,
        tp.content,
        geo.puma_names_from_geos(geo.parse_acs_geos(real(reqs["acs_geos"]).content)),
        None,
    )[CBSA]
    chosen = sorted(p.puma for p in g.pumas if p.full and p.state == PUMS_STATE)[:N_PUMAS]
    h_bytes, p_bytes, pums_rule = cut_pums(
        real(reqs[f"pums_h:{PUMS_STATE}"]).content,
        real(reqs[f"pums_p:{PUMS_STATE}"]).content,
        chosen,
        cfg.pums.housing_vars,
        cfg.pums.person_vars,
    )
    save(f"pums_h:{PUMS_STATE}", h_bytes, pums_rule)
    save(f"pums_p:{PUMS_STATE}", p_bytes, pums_rule)
    oews_bytes, _ = cut_oews(real(reqs["oews"]).content)
    save(
        "oews",
        oews_bytes,
        "AREA = 16980：总计与大类行、按职业代码排序的前 60 个细类、带 # / * / ** 标记的前 10 个细类",
    )
    save("ep_education", real(reqs["ep_education"]).content, "原文件（118 KB，不切）")
    save("cpi", cut_cpi(real(reqs["cpi"]).content), "CUUR0000SA0 在 2022 年及以后的行")
    items = basket_items(cfg_doc) | set(AP_ENERGY)
    areas = {"0000", cfg.metros[0].census_region, cfg.metros[0].ap_metro_area}
    s_bytes, d_bytes = cut_ap(
        real(reqs["ap:ap.series"]).content, real(reqs["ap:ap.data.0.Current"]).content, items, areas
    )
    save("ap:ap.series", s_bytes, f"地区 {sorted(areas)} × 食品篮子用到的项目与能源项目")
    save("ap:ap.data.0.Current", d_bytes, "上述序列在 2023 年及以后的行")
    for name in ("ap:ap.item", "ap:ap.area"):
        save(name, real(reqs[name]).content, "原文件（很小，不切）")
    for name in [n for n in reqs if n.startswith("ce:")]:
        save(name, real(reqs[name]).content, "原文件（很小，不切）")
    hd_bytes, py_bytes, _, cips = cut_ipeds(
        real(reqs["ipeds_hd"]).content, real(reqs["ipeds_ic_py"]).content
    )
    save(
        "ipeds_hd",
        hd_bytes,
        f"CBSA = 16980 且有按项目报告学费的院校中编号最小的 {N_INSTITUTIONS} 所",
    )
    save("ipeds_ic_py", py_bytes, "同上院校的项目行")
    save("cip_soc", cut_cip_soc(real(reqs["cip_soc"]).content, cips), "上述项目的 CIP 代码的对照行")
    for name in ("aqi", "aqi_2024"):
        save(name, cut_aqi(real(reqs[name]).content), "表头 + CBSA Code = 16980")
    save(
        "bea_rpp",
        cut_rpp(real(reqs["bea_rpp"]).content),
        "MARPP_MSA 文件里 GeoFIPS = 00000（全美）与 16980 的行",
    )
    for t in sources.ACS_TABLES:
        name = f"acs_table:{t}"
        save(
            name,
            cut_summary_table(real(reqs[name]).content, cbsa_geo_id),
            f"表头 + GEO_ID = {cbsa_geo_id}",
        )
    for cat in sources.ENERGYSTAR:
        name = f"energystar:{cat}"
        rows = json.loads(real(reqs[name]).content)
        save(name, json.dumps(rows[:6], ensure_ascii=False).encode("utf-8"), "文件里的前 6 个型号")
    chi_bytes, chi_rule = cut_chicago(real(reqs["chicago_activities"]).content)
    save("chicago_activities", chi_bytes, chi_rule)
    save(
        f"osm_food:{CBSA}",
        cut_osm(real(reqs[f"osm_food:{CBSA}"]).content, "food"),
        "前 20 个带标签的元素 + 写了外带 / 外送 / 得来速的前 10 个",
    )
    save(
        f"osm_leisure:{CBSA}",
        cut_osm(real(reqs[f"osm_leisure:{CBSA}"]).content, "leisure"),
        "前 30 个带标签的元素 + 私人场所的前 3 个（检验过滤）",
    )
    live = json.loads((cache / "_fetch_log.json").read_text(encoding="utf-8"))
    (out_cache / "_fetch_log.json").write_text(
        json.dumps({"hud_fmr": live["hud_fmr"]}, ensure_ascii=False, indent=1, sort_keys=True)
        + "\n",
        encoding="utf-8",
        newline="\n",
    )
    total = sum(sizes.values())
    for name, n in sorted(sizes.items(), key=lambda kv: -kv[1]):
        print(f"{n:>9}  {name}")
    print(f"共 {len(sizes)} 份样本，{total} 字节")
    return 0


if __name__ == "__main__":
    import sys

    sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    raise SystemExit(main())
